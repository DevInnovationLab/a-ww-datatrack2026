"""Render sessionN_slide_content.md as a PowerPoint deck in the Session 4 style.

Usage:  python3 presentations/build_pptx.py 5
Output: presentations/session5_slides.pptx
Needs:  python-pptx (pip install python-pptx)

The design copies presentations/session4_slides.pptx, which was composed by
hand: Montserrat headings, Arial body, maroon (8B0021 / 8C1038), gold
(FAA61A) and charcoal (303540). That deck is also the base file, so its theme
and slide size carry over; its slides are dropped. Speaker notes come from the
"> **Notes:**" blocks. Re-run after editing the markdown.

Slide types, from the "## Slide N — <title> · <timing>" headings (the same
markdown that build_slides.py turns into the HTML review deck):
  Title                     maroon title slide
  Why we're here            light statement slide
  Section divider           charcoal divider with the part number and timing
  EXERCISE ... / SPOT ...   "YOUR TURN" slide; the last bold line is a banner
  anything else             content slide
Content blocks are laid out top to bottom: italic lines become the lead,
"**Label**" + bullets (two or more groups) become cards, bullets and numbered
lists become the body, fenced code becomes a code box, a closing bold line
becomes the takeaway box, and the column markers become two columns.
"""

import math
import re
import sys
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import build_slides as md  # noqa: E402  (shares the markdown parser)

HERE = Path(__file__).parent
BASE = HERE / "session4_slides.pptx"

MAROON, MAROON2, GOLD = "8B0021", "8C1038", "FAA61A"
CHARCOAL, INK, GREY, CARD = "303540", "1A1A1A", "6B6E76", "F3F3F3"
HEAD, BODY, MONO = "Montserrat", "Arial", "Courier New"
FOOTER = "Development Innovation Lab / University of Chicago"

SESSIONS = {5: dict(kicker="FROM ANALYSIS TO PUBLICATION")}

L, W = 0.7, 11.9          # left margin and content width (inches)
TOP, BOTTOM = 1.95, 6.85  # content area below the title, above the footer


# ---- text helpers ------------------------------------------------------------

def rgb(h):
    return RGBColor.from_string(h)


def runs(text):
    """Split inline markdown into (text, bold, italic, code) runs."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False, False))
        elif part.startswith("`"):
            out.append((part[1:-1], False, False, True))
        elif part.startswith("*"):
            out.append((part[1:-1], False, True, False))
        else:
            out.append((part, False, False, False))
    return out


def plain(text):
    return "".join(r[0] for r in runs(text))


def box(slide, x, y, w, h, name=None, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tb


def write(tf, paras, size, color, font=BODY, bold=False, bullet=None,
          space_before=0, line=1.0, align=PP_ALIGN.LEFT, first=True, literal=False):
    """paras: list of strings (inline markdown unless literal).
    bullet: None, '•', '☐' or 'num'."""
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        if i or not first:
            p.space_before = Pt(space_before)
        pPr = p._p.get_or_add_pPr()
        if bullet:
            pPr.set("marL", str(int(Inches(0.38))))
            pPr.set("indent", str(-int(Inches(0.38))))
            buclr = pPr.makeelement(qn("a:buClr"), {})
            clr = buclr.makeelement(qn("a:srgbClr"), {"val": MAROON if bullet == "num" else color})
            buclr.append(clr)
            pPr.append(buclr)
            if bullet == "num":
                pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": BODY}))
                pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
            else:
                pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": BODY}))
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": bullet}))
        for t, b, it, code in ([(text, False, False, False)] if literal else runs(text)):
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(size)
            f.bold = bold or b
            f.italic = it
            f.name = MONO if code else font
            f.color.rgb = rgb(color)


def text_height(paras, size, width, line=1.15, space=0, mono=False, indent=0.0):
    """Rough height in inches of wrapped paragraphs (Arial ~0.5 em per char)."""
    cpl = max(8, (width - indent) * 72 / (size * (0.6 if mono else 0.5)))
    lines = sum(max(1, math.ceil(len(plain(p)) / cpl)) for p in paras)
    return lines * size * line / 72 + max(0, len(paras) - 1) * space / 72


def rect(slide, x, y, w, h, color, name):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.name = name
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def background(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def chrome(slide, n, kicker, title, dark=False):
    k = box(slide, L, 0.5, 10.0, 0.4, "Kicker")
    write(k.text_frame, [kicker], 12, GOLD if dark else MAROON, HEAD, bold=True)
    if title:
        t = box(slide, L, 0.95, W, 0.85, "Title")
        write(t.text_frame, [title], 26, CHARCOAL, HEAD, bold=True)
    fc = "D8D3D5" if dark else CHARCOAL
    f = box(slide, 0.5, 7.05, 9.5, 0.3, "Footer")
    write(f.text_frame, [FOOTER], 9, fc)
    s = box(slide, 12.33, 7.05, 0.5, 0.3, "Slide number")
    write(s.text_frame, [str(n)], 9, fc, align=PP_ALIGN.RIGHT)


def notes(slide, text):
    text = re.sub(r"\*\*|`", "", text).strip()
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ---- markdown content -> blocks ----------------------------------------------

def blocks(content):
    """Turn slide markdown into a list of (kind, payload) blocks."""
    out, lines, i = [], content.split("\n"), 0
    while i < len(lines):
        s = lines[i].rstrip()
        st = s.strip()
        if not st:
            i += 1
            continue
        if st.startswith("```"):
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            out.append(("code", lines[i + 1:j]))
            i = j + 1
            continue
        if st == "<!-- columns -->":
            j = lines.index("<!-- end columns -->", i)
            seg = "\n".join(lines[i + 1:j]).split("<!-- next column -->")
            out.append(("columns", [blocks(c) for c in seg]))
            i = j + 1
            continue
        m = re.match(r"^(\s*)(- \[ \] |- |\d+\. )(.*)$", s)
        if m:
            kind = "num" if m.group(2)[0].isdigit() else ("check" if "[ ]" in m.group(2) else "bullet")
            items = []
            while i < len(lines):
                m = re.match(r"^(\s*)(- \[ \] |- |\d+\. )(.*)$", lines[i])
                if m:
                    items.append(m.group(3))
                elif lines[i].startswith("   ") and lines[i].strip() and items:
                    items[-1] += " · " + lines[i].strip()   # continuation line
                else:
                    break
                i += 1
            out.append((kind, items))
            continue
        if re.match(r"^\*[^*].*\*$", st):
            out.append(("lead", st[1:-1]))
        elif re.match(r"^\*\*[^*]+\*\*$", st):
            out.append(("bold", st[2:-2]))
        else:
            out.append(("para", st))
        i += 1
    # "**Label**" followed by a list, twice or more -> cards
    groups, k = [], 0
    while k < len(out) - 1 and out[k][0] == "bold" and out[k + 1][0] in ("bullet", "num", "check"):
        groups.append((out[k][1], out[k + 1][1]))
        k += 2
    if len(groups) >= 2:
        out = [("cards", groups)] + out[k:]
    return out


# ---- slide builders ----------------------------------------------------------

def title_slide(prs, content, nt):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    background(s, MAROON)
    lines = [l for l in content.split("\n") if l.strip()]
    kicker = next(l for l in lines if l.startswith("**")).strip("*")
    title = next(l for l in lines if l.startswith("# "))[2:]
    sub = next((l for l in lines if l.startswith("*") and not l.startswith("**")), "")
    rest = [l.strip() for l in lines if l not in (sub,) and not l.startswith(("#", "**"))]
    people = [p.strip() for p in rest[0].split(" · ")] if rest else []
    write(box(s, L, 0.5, 11.0, 0.4, "Kicker").text_frame, [kicker], 12, GOLD, HEAD, bold=True)
    write(box(s, L, 2.1, 11.5, 2.0, "Title", MSO_ANCHOR.BOTTOM).text_frame, [title], 40, "FFFFFF", HEAD, bold=True)
    write(box(s, L, 4.3, 10.5, 0.7, "Subtitle").text_frame, [sub.strip("*")], 16, "F1DDE2")
    write(box(s, L, 5.25, 10.0, 1.0, "Authors").text_frame, people, 13, GOLD, bold=True)
    if len(rest) > 1:
        write(box(s, L, 6.5, 10.0, 0.4, "Affiliation").text_frame, rest[1:], 12, "F1DDE2")
    notes(s, nt)


def statement_slide(prs, title, content, nt, bg_xml):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    if bg_xml is not None:
        s._element.cSld.insert(0, deepcopy(bg_xml))
    lines = [plain(l.lstrip("- ")) for l in content.split("\n") if l.strip()]
    write(box(s, L, 2.6, W, 1.0, "Question", MSO_ANCHOR.BOTTOM).text_frame, [title], 40, MAROON, HEAD, bold=True)
    write(box(s, L, 3.85, W, 1.2, "Subtitle").text_frame, [" · ".join(lines)], 18, GREY)
    notes(s, nt)


def divider_slide(prs, n, content, nt, minutes, kicker):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    background(s, CHARCOAL)
    lines = [l for l in content.split("\n") if l.strip() and not l.strip().startswith("~")]
    num, _, name = lines[0].strip("*").partition(" · ")
    write(box(s, L, 0.5, 10.0, 0.4, "Kicker").text_frame, [kicker], 12, GOLD, HEAD, bold=True)
    write(box(s, L, 1.6, 3.0, 1.5, "Section number", MSO_ANCHOR.BOTTOM).text_frame, [num], 60, GOLD, HEAD, bold=True)
    write(box(s, L, 3.05, 11.5, 1.3, "Title").text_frame, [name], 32, "FFFFFF", HEAD, bold=True)
    if len(lines) > 1:
        write(box(s, L, 4.35, 10.3, 1.0, "Subtitle").text_frame, lines[1:], 15, "E4E4E7")
    if minutes:
        write(box(s, L, 6.5, 6.0, 0.4, "Timing").text_frame, [f"~{minutes:g} minutes"], 12, GOLD)
    f = box(s, 0.5, 7.05, 9.5, 0.3, "Footer")
    write(f.text_frame, [FOOTER], 9, "D8D3D5")
    sn = box(s, 12.33, 7.05, 0.5, 0.3, "Slide number")
    write(sn.text_frame, [str(n)], 9, "D8D3D5", align=PP_ALIGN.RIGHT)
    notes(s, nt)


def place(slide, bl, x, y, w, ymax, exercise, scale=1.0):
    """Lay blocks out from y down; return the y reached (dry run if slide is None)."""
    def sz(v):
        return max(12, round(v * scale))

    for kind, pay in bl:
        if kind == "lead":
            h = text_height([pay], 15, w)
            if slide:
                write(box(slide, x, y, w, h + 0.05, "Lead").text_frame, [pay], 15, CHARCOAL)
            y += h + 0.15
        elif kind == "bold" and kind == bl[-1][0] and pay == bl[-1][1] and len(bl) > 1:
            fs = sz(19)
            h = max(0.75, text_height([pay], fs, w - 0.6) + 0.35)
            y = max(y + 0.15, min(y + 0.3, ymax - h))
            if slide:
                rect(slide, x, y, w, h, MAROON2 if exercise else CARD, "Banner" if exercise else "Takeaway")
                write(box(slide, x + 0.3, y, w - 0.6, h, "Takeaway text", MSO_ANCHOR.MIDDLE).text_frame,
                      [pay], fs, "FFFFFF" if exercise else CHARCOAL, bold=True)
            y += h + 0.2
        elif kind == "bold":
            if "→" in pay and len(pay) < 60 and w > 10:
                steps = [p.strip() for p in pay.split("→")]
                bw = (w - 0.55 * (len(steps) - 1)) / len(steps)
                if slide:
                    for k, st in enumerate(steps):
                        bx = x + k * (bw + 0.55)
                        last = k == len(steps) - 1
                        rect(slide, bx, y, bw, 0.75, MAROON2 if last else CARD, f"Flow {st}")
                        write(box(slide, bx, y, bw, 0.75, None, MSO_ANCHOR.MIDDLE).text_frame, [st], 18,
                              "FFFFFF" if last else CHARCOAL, HEAD, bold=True, align=PP_ALIGN.CENTER)
                        if not last:
                            write(box(slide, bx + bw, y, 0.55, 0.75, None, MSO_ANCHOR.MIDDLE).text_frame,
                                  ["→"], 24, MAROON, bold=True, align=PP_ALIGN.CENTER)
                y += 1.05
            else:
                fs = sz(18)
                h = text_height([pay], fs, w)
                if slide:
                    write(box(slide, x, y, w, h + 0.05, "Subhead").text_frame, [pay], fs, CHARCOAL, bold=True)
                y += h + 0.2
        elif kind in ("bullet", "num", "check"):
            fs = sz(19)
            bu = {"bullet": "•", "num": "num", "check": "☐"}[kind]
            h = text_height(pay, fs, w, space=10, indent=0.38)
            if slide:
                write(box(slide, x, y, w, h + 0.1, "Body").text_frame, pay, fs, INK,
                      bullet=bu, space_before=10, line=1.1)
            y += h + 0.25
        elif kind == "para":
            fs = sz(15)
            h = text_height([pay], fs, w)
            if slide:
                write(box(slide, x, y, w, h + 0.05, "Text").text_frame, [pay], fs, CHARCOAL)
            y += h + 0.2
        elif kind == "code":
            fs = sz(15) if w > 8 else sz(13)
            h = text_height(pay, fs, w - 0.6, line=1.1, mono=True) + 0.4
            if slide:
                rect(slide, x, y, w, h, CARD, "Code")
                tb = box(slide, x + 0.3, y + 0.2, w - 0.6, h - 0.4, "Code text")
                write(tb.text_frame, pay or [""], fs, CHARCOAL, MONO, line=1.0, literal=True)
            y += h + 0.2
        elif kind == "cards":
            n = len(pay)
            gap = 0.35
            cw = (w - gap * (n - 1)) / n
            fs = sz(17)
            hs = [0.45 + text_height(items, fs, cw - 0.56, space=8, indent=0.38) for _, items in pay]
            ch = min(max(hs) + 0.4, ymax - y - 1.0) if bl[-1][0] == "bold" else max(hs) + 0.4
            ch = max(ch, max(hs) + 0.3)
            if slide:
                for k, (lab, items) in enumerate(pay):
                    cx = x + k * (cw + gap)
                    rect(slide, cx, y, cw, ch, CARD, "Card")
                    write(box(slide, cx + 0.28, y + 0.22, cw - 0.56, 0.35, "Card label").text_frame,
                          [plain(lab).upper()], 11, MAROON, HEAD, bold=True)
                    write(box(slide, cx + 0.28, y + 0.67, cw - 0.56, ch - 0.8, "Card text").text_frame,
                          items, fs, INK, bullet="•", space_before=8, line=1.1)
            y += ch + 0.3
        elif kind == "columns":
            n = len(pay)
            gap = 0.5
            cw = (w - gap * (n - 1)) / n
            ys = []
            for k, col in enumerate(pay):
                ys.append(place(slide, col, x + k * (cw + gap), y, cw, ymax, exercise, scale))
            y = max(ys)
    return y


def content_slide(prs, n, kicker, title, content, nt, exercise):
    bl = blocks(content)
    scale = 1.0
    while scale > 0.7 and place(None, bl, L, TOP, W, BOTTOM, exercise, scale) > BOTTOM:
        scale -= 0.05
    s = prs.slides.add_slide(prs.slide_layouts[0])
    background(s, "FFFFFF")
    chrome(s, n, kicker, title)
    place(s, bl, L, TOP, W, BOTTOM, exercise, scale)
    notes(s, nt)


# ---- deck --------------------------------------------------------------------

def exercise_title(title):
    head, _, rest = title.partition(" · ")
    if not rest:
        return title
    head = head.capitalize()
    return f"{head}: {rest[0].lower() + rest[1:]}"


def build(session):
    cfg = SESSIONS[session]
    src = HERE / f"session{session}_slide_content.md"
    out = HERE / f"session{session}_slides.pptx"
    _, slides = md.split_slides(src.read_text())

    prs = Presentation(str(BASE))
    bg_xml = prs.slides[1]._element.cSld.find(qn("p:bg")) if len(prs.slides) > 1 else None
    bg_xml = deepcopy(bg_xml) if bg_xml is not None else None
    sld_ids = prs.slides._sldIdLst
    for sid in list(sld_ids):
        prs.part.drop_rel(sid.get(qn("r:id")))
        sld_ids.remove(sid)

    # minutes per section, for the divider timing line
    parsed = []
    for heading, body in slides:
        label, title, timing, flags = md.parse_heading(heading)
        content, nts = md.split_notes(body)
        parsed.append((title, md.minutes(timing), content, nts))
    sec_min, cur = {}, None
    for k, (title, mins, _, _) in enumerate(parsed):
        if title == "Section divider":
            cur = k
            sec_min[cur] = mins
        elif cur is not None:
            sec_min[cur] += mins

    part, part_name = None, cfg["kicker"]
    for k, (title, mins, content, nts) in enumerate(parsed):
        n = k + 1
        if title == "Title":
            title_slide(prs, content, nts)
        elif title == "Section divider":
            first = [l for l in content.split("\n") if l.strip()][0].strip("*")
            num, _, name = first.partition(" · ")
            part, part_name = int(num), name.upper()
            divider_slide(prs, n, content, nts, sec_min[k], cfg["kicker"])
        elif title.lower().startswith("why we"):
            statement_slide(prs, title, content, nts, bg_xml)
        else:
            exercise = bool(re.match(r"^(EXERCISE|SPOT THE)", title))
            if exercise:
                kick = f"PART {part} · YOUR TURN" if part else "YOUR TURN"
                title = exercise_title(title)
                content = re.sub(r"^\*(.+)\*$", lambda m: f"*{m.group(1)} · {mins:g} minutes*" if mins else m.group(0),
                                 content, count=1, flags=re.M)
            else:
                kick = f"PART {part} · {part_name}" if part else cfg["kicker"]
            content_slide(prs, n, kick, title, content, nts, exercise)

    prs.save(str(out))
    print(f"Wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build(int(sys.argv[1]) if len(sys.argv) > 1 else 5)
